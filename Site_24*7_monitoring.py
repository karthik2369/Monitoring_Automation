"""
Site24x7 Disk Report Automation — Phase 1
==========================================
Flow:
    Site24x7 India API
        -> Python
        -> disk metrics for selected servers
        -> Excel report
        -> optional Outlook email

Phase 1 deliberately does NOT include Jenkins, Grafana, InfluxDB,
or browser automation. Those can be added in later phases.

Install:
    pip install requests openpyxl

Commands:
    python disk_report_phase1.py --validate
    python disk_report_phase1.py
    python disk_report_phase1.py --schedule

SECURITY:
    Keep credentials outside this file.
    Recommended: environment variables.

Site24x7 currently documents OAuth 2.0 for API access. For a
long-running automation, use a Zoho OAuth refresh token rather
than hard-coding a short-lived access token.
"""

import argparse
import html
import logging
import os
import smtplib
import ssl
import sys
import time
from datetime import datetime
from email.mime.base import MIMEBase
from email.mime.multipart import MIMEMultipart
from email.mime.text import MIMEText
from email import encoders
from pathlib import Path
from typing import Any, Dict, List, Optional

import openpyxl
import requests
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter


# ============================================================
# CONFIGURATION
# ============================================================

SITE24X7_BASE = os.getenv(
    "SITE24X7_BASE_URL",
    "https://www.site24x7.in/api"
)

# OAuth 2.0
ZOHO_ACCOUNTS_BASE = os.getenv(
    "ZOHO_ACCOUNTS_BASE_URL",
    "https://accounts.zoho.in"
)

# Preferred for automation:
ZOHO_CLIENT_ID = os.getenv("SITE24X7_CLIENT_ID", "")
ZOHO_CLIENT_SECRET = os.getenv("SITE24X7_CLIENT_SECRET", "")
ZOHO_REFRESH_TOKEN = os.getenv("SITE24X7_REFRESH_TOKEN", "")

# Optional: useful for a quick one-time test.
# Do NOT put this directly in the source file.
SITE24X7_ACCESS_TOKEN = os.getenv("SITE24X7_ACCESS_TOKEN", "")

# Target servers. Comma-separated:
# export SITE24X7_TARGET_SERVERS="server1,server2,server3"
TARGET_SERVERS = [
    s.strip()
    for s in os.getenv("SITE24X7_TARGET_SERVERS", "").split(",")
    if s.strip()
]

# Email is optional in Phase 1.
EMAIL_ENABLED = os.getenv("EMAIL_ENABLED", "false").lower() == "true"
EMAIL_FROM = os.getenv("EMAIL_FROM", "")
EMAIL_PASSWORD = os.getenv("EMAIL_PASSWORD", "")
EMAIL_TO = os.getenv("EMAIL_TO", "")
EMAIL_CC = [
    s.strip()
    for s in os.getenv("EMAIL_CC", "").split(",")
    if s.strip()
]

SMTP_SERVER = os.getenv("SMTP_SERVER", "smtp.office365.com")
SMTP_PORT = int(os.getenv("SMTP_PORT", "587"))

REPORT_DIR = Path(os.getenv("REPORT_DIR", "."))
REPORT_DIR.mkdir(parents=True, exist_ok=True)

WARN_THRESHOLD = float(os.getenv("DISK_WARN_THRESHOLD", "80"))
CRITICAL_THRESHOLD = float(os.getenv("DISK_CRITICAL_THRESHOLD", "90"))

REQUEST_TIMEOUT = 30
MAX_RETRIES = 3

STATUS_MAP = {
    0: "DOWN",
    1: "UP",
    2: "TROUBLE",
    5: "SUSPENDED",
    7: "MAINTENANCE",
    10: "CONFIG ERROR",
}


# ============================================================
# LOGGING
# ============================================================

logging.basicConfig(
    level=logging.INFO,
    format="[%(levelname)s] %(message)s",
)
logger = logging.getLogger("site24x7-disk-report")


# ============================================================
# HTTP / AUTH
# ============================================================

session = requests.Session()
session.headers.update({
    "Accept": "application/json; version=2.0",
})


def get_access_token() -> str:
    """
    Return an access token.

    Priority:
      1. Existing SITE24X7_ACCESS_TOKEN
      2. Refresh OAuth access token using:
         SITE24X7_CLIENT_ID
         SITE24X7_CLIENT_SECRET
         SITE24X7_REFRESH_TOKEN
    """
    if SITE24X7_ACCESS_TOKEN:
        return SITE24X7_ACCESS_TOKEN

    if not all([
        ZOHO_CLIENT_ID,
        ZOHO_CLIENT_SECRET,
        ZOHO_REFRESH_TOKEN,
    ]):
        raise RuntimeError(
            "No Site24x7 OAuth credentials configured. Set either "
            "SITE24X7_ACCESS_TOKEN or SITE24X7_CLIENT_ID, "
            "SITE24X7_CLIENT_SECRET and SITE24X7_REFRESH_TOKEN."
        )

    token_url = f"{ZOHO_ACCOUNTS_BASE}/oauth/v2/token"

    response = requests.post(
        token_url,
        data={
            "refresh_token": ZOHO_REFRESH_TOKEN,
            "client_id": ZOHO_CLIENT_ID,
            "client_secret": ZOHO_CLIENT_SECRET,
            "grant_type": "refresh_token",
        },
        timeout=REQUEST_TIMEOUT,
    )

    if response.status_code != 200:
        raise RuntimeError(
            f"OAuth token refresh failed: HTTP {response.status_code} "
            f"{response.text[:300]}"
        )

    body = response.json()
    access_token = body.get("access_token")

    if not access_token:
        raise RuntimeError("OAuth response did not contain access_token.")

    return access_token


def api_get(path: str, params: Optional[Dict[str, Any]] = None) -> Dict[str, Any]:
    """GET a Site24x7 API endpoint with retries."""
    token = get_access_token()

    headers = {
        "Authorization": f"Zoho-oauthtoken {token}",
    }

    url = f"{SITE24X7_BASE.rstrip('/')}/{path.lstrip('/')}"

    last_error = None

    for attempt in range(1, MAX_RETRIES + 1):
        try:
            response = session.get(
                url,
                headers=headers,
                params=params,
                timeout=REQUEST_TIMEOUT,
            )

            if response.status_code == 401:
                raise RuntimeError(
                    "Site24x7 returned HTTP 401. The access token is invalid "
                    "or expired. If using OAuth refresh-token mode, verify "
                    "client ID, client secret and refresh token."
                )

            response.raise_for_status()

            body = response.json()

            if body.get("code") not in (None, 0):
                raise RuntimeError(
                    f"Site24x7 API error {body.get('code')}: "
                    f"{body.get('message', 'Unknown error')}"
                )

            return body

        except (requests.RequestException, ValueError, RuntimeError) as exc:
            last_error = exc
            if attempt < MAX_RETRIES:
                wait = attempt * 2
                logger.warning(
                    "API request failed (attempt %s/%s): %s. Retrying in %ss...",
                    attempt, MAX_RETRIES, exc, wait
                )
                time.sleep(wait)
            else:
                break

    raise RuntimeError(f"Site24x7 API request failed: {last_error}")


# ============================================================
# SITE24X7 DATA
# ============================================================

def fetch_all_monitors() -> List[Dict[str, Any]]:
    """Fetch current monitor status data."""
    body = api_get("/current_status")
    data = body.get("data") or {}
    monitors = data.get("monitors") or []

    if not isinstance(monitors, list):
        raise RuntimeError("Unexpected Site24x7 response: monitors is not a list.")

    return monitors


def match_target_servers(monitors: List[Dict[str, Any]]) -> List[Dict[str, Any]]:
    """
    Match only configured server monitors.

    If no TARGET_SERVERS are configured, return all monitors whose
    monitor_type is SERVER. This is useful for initial testing.
    """
    server_monitors = [
        m for m in monitors
        if str(m.get("monitor_type", "")).upper() == "SERVER"
    ]

    if not TARGET_SERVERS:
        logger.warning(
            "SITE24X7_TARGET_SERVERS is empty. Using ALL server monitors "
            "for this test run."
        )
        return server_monitors

    wanted = {name.casefold() for name in TARGET_SERVERS}

    matched = [
        m for m in server_monitors
        if str(m.get("name", "")).casefold() in wanted
    ]

    found_names = {
        str(m.get("name", "")).casefold()
        for m in matched
    }

    missing = [
        name for name in TARGET_SERVERS
        if name.casefold() not in found_names
    ]

    if missing:
        logger.warning("Configured servers not found in Site24x7:")
        for name in missing:
            logger.warning("  - %s", name)

    return matched


def numeric(value: Any) -> Optional[float]:
    """Convert values such as '82.4%' or '100 GB' to a number."""
    if value is None:
        return None

    text = str(value).strip()
    cleaned = ""

    for char in text:
        if char.isdigit() or char in ".-":
            cleaned += char

    try:
        return float(cleaned)
    except ValueError:
        return None


def extract_metric(
    attributes: Any,
    keywords: List[str],
) -> Any:
    """
    Best-effort extraction from current_status attributes.

    IMPORTANT:
    Site24x7 can expose different attribute structures depending on
    monitor type/configuration. The script intentionally does not
    pretend that a guessed field name is guaranteed.
    """
    if not isinstance(attributes, list):
        return None

    for attr in attributes:
        if not isinstance(attr, dict):
            continue

        name = str(
            attr.get("name")
            or attr.get("display_name")
            or attr.get("attribute")
            or ""
        ).casefold()

        if any(keyword.casefold() in name for keyword in keywords):
            return attr.get("value")

    return None


def extract_disk_records(
    monitors: List[Dict[str, Any]]
) -> List[Dict[str, Any]]:
    """
    Extract server-level disk metrics from current_status.

    The output is intentionally one row per server, matching the
    Phase 1 requirement supplied in the project.

    If your existing Excel requires one row per C:/D:/E: partition,
    we should switch this function to the partition-level API/report
    structure after inspecting one real API response.
    """
    records = []
    report_time = datetime.now().strftime("%Y-%m-%d %H:%M:%S")

    for monitor in monitors:
        name = monitor.get("name", "N/A")
        hostname = monitor.get("hostname") or name

        status_code = monitor.get("status", -1)
        try:
            status_code = int(status_code)
        except (TypeError, ValueError):
            status_code = -1

        status = STATUS_MAP.get(
            status_code,
            f"UNKNOWN ({status_code})"
        )

        attributes = monitor.get("attributes", [])

        free_value = extract_metric(
            attributes,
            [
                "overall disk free",
                "disk free",
                "free disk",
            ],
        )

        used_value = extract_metric(
            attributes,
            [
                "overall disk used",
                "disk utilization",
                "disk usage",
                "disk used",
                "used disk",
            ],
        )

        # Keep the raw value rather than inventing a unit.
        free_disk = str(free_value) if free_value is not None else "N/A"

        used_numeric = numeric(used_value)
        used_disk = (
            f"{used_numeric:.2f}%"
            if used_numeric is not None
            else "N/A"
        )

        records.append({
            "name": name,
            "hostname": hostname,
            "status": status,
            "free_disk": free_disk,
            "used_disk": used_disk,
            "last_polled": monitor.get("last_polled_time", "N/A"),
            "report_time": report_time,
        })

    return records


# ============================================================
# EXCEL
# ============================================================

def disk_fill(value: str):
    pct = numeric(value)

    if pct is None:
        return PatternFill(fill_type=None)

    if pct >= CRITICAL_THRESHOLD:
        return PatternFill("solid", fgColor="FFC7CE")

    if pct >= WARN_THRESHOLD:
        return PatternFill("solid", fgColor="FFD966")

    return PatternFill(fill_type=None)


def generate_excel(records: List[Dict[str, Any]]) -> Path:
    """Create the Phase 1 Excel report."""
    today = datetime.now().strftime("%Y-%m-%d")
    output = REPORT_DIR / f"Disk_Report_{today}.xlsx"

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Disk Report"

    header_fill = PatternFill("solid", fgColor="1F4E79")
    alt_fill = PatternFill("solid", fgColor="DCE6F1")
    up_fill = PatternFill("solid", fgColor="C6EFCE")
    down_fill = PatternFill("solid", fgColor="FFC7CE")
    trouble_fill = PatternFill("solid", fgColor="FFEB9C")

    title_font = Font(
        name="Calibri",
        bold=True,
        size=13,
        color="1F4E79",
    )
    header_font = Font(
        name="Calibri",
        bold=True,
        size=11,
        color="FFFFFF",
    )
    data_font = Font(name="Calibri", size=10)

    border = Border(
        left=Side(style="thin"),
        right=Side(style="thin"),
        top=Side(style="thin"),
        bottom=Side(style="thin"),
    )

    center = Alignment(horizontal="center", vertical="center")
    left = Alignment(horizontal="left", vertical="center")

    generated = datetime.now().strftime("%d %B %Y, %H:%M:%S")

    ws.merge_cells("A1:G1")
    ws["A1"] = f"Site24x7 Disk Space Report | {generated}"
    ws["A1"].font = title_font
    ws["A1"].alignment = center
    ws.row_dimensions[1].height = 30

    columns = [
        "Monitor Name",
        "Hostname",
        "Status",
        "Free Disk Space",
        "Used Disk %",
        "Last Polled",
        "Report Time",
    ]

    for col_idx, name in enumerate(columns, 1):
        cell = ws.cell(row=2, column=col_idx, value=name)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = center
        cell.border = border

    for row_idx, record in enumerate(records, start=3):
        values = [
            record["name"],
            record["hostname"],
            record["status"],
            record["free_disk"],
            record["used_disk"],
            record["last_polled"],
            record["report_time"],
        ]

        for col_idx, value in enumerate(values, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=value)
            cell.font = data_font
            cell.border = border

            if col_idx == 3:
                cell.alignment = center
                if value == "UP":
                    cell.fill = up_fill
                elif value == "DOWN":
                    cell.fill = down_fill
                elif value not in ("UP", "N/A"):
                    cell.fill = trouble_fill

            elif col_idx == 5:
                cell.alignment = center
                cell.fill = disk_fill(str(value))

            else:
                cell.alignment = left
                if row_idx % 2 == 0:
                    cell.fill = alt_fill

    widths = [28, 24, 16, 22, 16, 24, 22]
    for idx, width in enumerate(widths, 1):
        ws.column_dimensions[get_column_letter(idx)].width = width

    ws.freeze_panes = "A3"
    ws.auto_filter.ref = f"A2:G{max(2, len(records) + 2)}"

    summary_row = len(records) + 5

    up = sum(1 for r in records if r["status"] == "UP")
    down = sum(1 for r in records if r["status"] == "DOWN")
    other = len(records) - up - down

    summary = [
        ("Total Servers", len(records)),
        ("UP", up),
        ("DOWN", down),
        ("TROUBLE / OTHER", other),
    ]

    for offset, (label, value) in enumerate(summary):
        ws.cell(row=summary_row + offset, column=1, value=label).font = Font(
            bold=True
        )
        ws.cell(row=summary_row + offset, column=2, value=value)

    # Save safely through a temporary file so a failed write does not
    # destroy the previous report.
    temp_output = output.with_suffix(".tmp.xlsx")
    wb.save(temp_output)
    temp_output.replace(output)

    logger.info("Excel report saved: %s", output)
    return output


# ============================================================
# EMAIL
# ============================================================

def status_bg(status: str) -> str:
    return {
        "UP": "#C6EFCE",
        "DOWN": "#FFC7CE",
        "TROUBLE": "#FFEB9C",
        "CONFIG ERROR": "#FFEB9C",
    }.get(status, "#F2F2F2")


def disk_bg(value: str) -> str:
    pct = numeric(value)

    if pct is None:
        return "#FFFFFF"

    if pct >= CRITICAL_THRESHOLD:
        return "#FFC7CE"

    if pct >= WARN_THRESHOLD:
        return "#FFD966"

    return "#FFFFFF"


def build_email_html(records: List[Dict[str, Any]]) -> str:
    up = sum(1 for r in records if r["status"] == "UP")
    down = sum(1 for r in records if r["status"] == "DOWN")
    other = len(records) - up - down

    rows = []

    for record in records:
        rows.append(
            f"""
            <tr>
              <td>{html.escape(str(record["name"]))}</td>
              <td>{html.escape(str(record["hostname"]))}</td>
              <td style="background:{status_bg(record["status"])};text-align:center;font-weight:bold">
                {html.escape(str(record["status"]))}
              </td>
              <td style="text-align:center">{html.escape(str(record["free_disk"]))}</td>
              <td style="background:{disk_bg(record["used_disk"])};text-align:center">
                {html.escape(str(record["used_disk"]))}
              </td>
              <td>{html.escape(str(record["last_polled"]))}</td>
            </tr>
            """
        )

    generated = datetime.now().strftime("%d %B %Y, %H:%M:%S")

    return f"""
    <html>
    <body style="font-family:Arial,sans-serif;color:#333">
      <h2 style="color:#1F4E79">Daily Disk Space Report</h2>

      <p>
        Generated: <strong>{html.escape(generated)}</strong><br>
        Servers monitored: <strong>{len(records)}</strong>
      </p>

      <p>
        <strong>UP:</strong> {up}
        &nbsp;&nbsp;
        <strong>DOWN:</strong> {down}
        &nbsp;&nbsp;
        <strong>TROUBLE/OTHER:</strong> {other}
      </p>

      <table style="border-collapse:collapse;width:100%">
        <thead>
          <tr style="background:#1F4E79;color:white">
            <th style="padding:8px;border:1px solid #ccc">Monitor Name</th>
            <th style="padding:8px;border:1px solid #ccc">Hostname</th>
            <th style="padding:8px;border:1px solid #ccc">Status</th>
            <th style="padding:8px;border:1px solid #ccc">Free Disk</th>
            <th style="padding:8px;border:1px solid #ccc">Used Disk %</th>
            <th style="padding:8px;border:1px solid #ccc">Last Polled</th>
          </tr>
        </thead>
        <tbody>
          {''.join(rows)}
        </tbody>
      </table>

      <p style="font-size:12px;color:#777">
        Yellow = disk usage &ge; {WARN_THRESHOLD:.0f}%.
        Red = disk usage &ge; {CRITICAL_THRESHOLD:.0f}% or server DOWN.
      </p>
    </body>
    </html>
    """


def send_email(excel_path: Path, records: List[Dict[str, Any]]) -> bool:
    if not EMAIL_ENABLED:
        logger.info("Email disabled. Set EMAIL_ENABLED=true to enable it.")
        return False

    missing = [
        name for name, value in [
            ("EMAIL_FROM", EMAIL_FROM),
            ("EMAIL_PASSWORD", EMAIL_PASSWORD),
            ("EMAIL_TO", EMAIL_TO),
        ]
        if not value
    ]

    if missing:
        logger.error(
            "Email is enabled but missing configuration: %s",
            ", ".join(missing),
        )
        return False

    message = MIMEMultipart("mixed")
    message["From"] = EMAIL_FROM
    message["To"] = EMAIL_TO

    if EMAIL_CC:
        message["CC"] = ", ".join(EMAIL_CC)

    message["Subject"] = (
        f"Daily Disk Space Report — "
        f"{datetime.now().strftime('%d %B %Y')}"
    )

    message.attach(MIMEText(
        build_email_html(records),
        "html",
        "utf-8",
    ))

    with open(excel_path, "rb") as file:
        attachment = MIMEBase(
            "application",
            "vnd.openxmlformats-officedocument.spreadsheetml.sheet",
        )
        attachment.set_payload(file.read())

    encoders.encode_base64(attachment)
    attachment.add_header(
        "Content-Disposition",
        f'attachment; filename="{excel_path.name}"',
    )
    message.attach(attachment)

    recipients = [EMAIL_TO] + EMAIL_CC

    try:
        context = ssl.create_default_context()

        with smtplib.SMTP(
            SMTP_SERVER,
            SMTP_PORT,
            timeout=REQUEST_TIMEOUT,
        ) as smtp:
            smtp.ehlo()
            smtp.starttls(context=context)
            smtp.ehlo()
            smtp.login(EMAIL_FROM, EMAIL_PASSWORD)
            smtp.sendmail(
                EMAIL_FROM,
                recipients,
                message.as_string(),
            )

        logger.info("Email sent to %s", EMAIL_TO)
        return True

    except smtplib.SMTPAuthenticationError:
        logger.error(
            "SMTP authentication failed. Your Microsoft 365 "
            "environment may require an approved authentication method."
        )
        return False

    except smtplib.SMTPException as exc:
        logger.error("SMTP error: %s", exc)
        return False


# ============================================================
# VALIDATION / DIAGNOSTICS
# ============================================================

def validate() -> bool:
    """Validate Site24x7 access and configured targets."""
    logger.info("Starting validation...")
    ok = True

    if not SITE24X7_ACCESS_TOKEN and not all([
        ZOHO_CLIENT_ID,
        ZOHO_CLIENT_SECRET,
        ZOHO_REFRESH_TOKEN,
    ]):
        logger.error(
            "No Site24x7 authentication configured."
        )
        ok = False
    else:
        try:
            monitors = fetch_all_monitors()
            servers = [
                m for m in monitors
                if str(m.get("monitor_type", "")).upper() == "SERVER"
            ]

            logger.info(
                "Site24x7 API connection successful: %d total monitors, %d server monitors.",
                len(monitors),
                len(servers),
            )

            matched = match_target_servers(monitors)

            if not matched:
                logger.error("No target server monitors matched.")
                ok = False
            else:
                logger.info(
                    "Target servers matched: %d",
                    len(matched),
                )

                # Diagnostic: show available attribute names for the first
                # matched server without exposing credentials.
                first = matched[0]
                attrs = first.get("attributes", [])

                if isinstance(attrs, list):
                    names = [
                        str(a.get("name"))
                        for a in attrs
                        if isinstance(a, dict) and a.get("name")
                    ]

                    logger.info(
                        "First matched server: %s",
                        first.get("name", "N/A"),
                    )
                    logger.info(
                        "Available current-status attributes: %s",
                        ", ".join(names[:30]) if names else "NONE",
                    )

        except Exception as exc:
            logger.error("Validation failed: %s", exc)
            ok = False

    if EMAIL_ENABLED:
        if not EMAIL_FROM or not EMAIL_PASSWORD or not EMAIL_TO:
            logger.error(
                "EMAIL_ENABLED=true but email configuration is incomplete."
            )
            ok = False
        else:
            logger.info(
                "Email configuration is present. SMTP login will be tested "
                "only when a report is sent."
            )

    if ok:
        logger.info("VALIDATION PASSED")
    else:
        logger.error("VALIDATION FAILED")

    return ok


# ============================================================
# MAIN
# ============================================================

def run_report() -> Optional[Path]:
    logger.info("=" * 60)
    logger.info("Site24x7 Disk Report — Phase 1")
    logger.info("=" * 60)

    monitors = fetch_all_monitors()

    logger.info(
        "Fetched %d monitors from Site24x7.",
        len(monitors),
    )

    matched = match_target_servers(monitors)

    if not matched:
        raise RuntimeError(
            "No matching server monitors found. "
            "Check SITE24X7_TARGET_SERVERS."
        )

    records = extract_disk_records(matched)

    # Guard against silently producing a useless report.
    missing_disk = [
        r["name"]
        for r in records
        if r["free_disk"] == "N/A"
        and r["used_disk"] == "N/A"
    ]

    if missing_disk:
        logger.warning(
            "No disk attributes were detected for: %s",
            ", ".join(missing_disk),
        )
        logger.warning(
            "Run --validate and inspect the listed current-status "
            "attributes before relying on the report."
        )

    excel_path = generate_excel(records)

    send_email(excel_path, records)

    logger.info("Report completed: %s", excel_path)

    return excel_path


def schedule_daily() -> None:
    """
    Simple standard-library scheduler.

    This is useful for testing Phase 1.
    For production, Phase 2 should run this script from Jenkins
    on a scheduled job instead of keeping Python alive forever.
    """
    logger.info("Daily scheduler started. Target time: 09:00.")

    while True:
        now = datetime.now()

        if now.hour == 9 and now.minute == 0:
            try:
                run_report()
            except Exception:
                logger.exception("Scheduled report failed.")

            # Prevent running twice in the same minute.
            time.sleep(65)

        time.sleep(20)


def parse_args():
    parser = argparse.ArgumentParser(
        description="Site24x7 Disk Report Automation — Phase 1"
    )

    parser.add_argument(
        "--validate",
        action="store_true",
        help="Validate Site24x7 authentication and target servers.",
    )

    parser.add_argument(
        "--schedule",
        action="store_true",
        help="Run the report every day at 09:00.",
    )

    return parser.parse_args()


def main() -> int:
    args = parse_args()

    try:
        if args.validate:
            return 0 if validate() else 1

        if args.schedule:
            schedule_daily()
            return 0

        run_report()
        return 0

    except KeyboardInterrupt:
        logger.info("Stopped.")
        return 130

    except Exception as exc:
        logger.error("Fatal error: %s", exc)
        return 1


if __name__ == "__main__":
    sys.exit(main())
